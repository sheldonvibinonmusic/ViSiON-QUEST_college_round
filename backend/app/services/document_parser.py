"""
Document parsing service for TXT, PDF, and DOCX files.
Preserves document identity, SHA-256 hash, and available page metadata.
"""
import io
from typing import List, Optional, Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument

from backend.app.schemas import SourceDocument
from backend.app.util.hashes import compute_sha256, compute_text_sha256, generate_id
from backend.app.errors import ErrorCode, SecondOpinionException

class ParsedPage:
    def __init__(self, page_number: Optional[int], text: str, start_char: int, end_char: int):
        self.page_number = page_number
        self.text = text
        self.start_char = start_char
        self.end_char = end_char

class ParsedResult:
    def __init__(self, document: SourceDocument, pages: List[ParsedPage], full_text: str):
        self.document = document
        self.pages = pages
        self.full_text = full_text

def parse_plain_text(text: str, filename: str = "pasted_source.txt") -> ParsedResult:
    """
    Parse raw string / TXT content.
    """
    clean_text = text.replace("\r\n", "\n").replace("\r", "\n")
    if not clean_text.strip():
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"Source document '{filename}' is empty or contains only whitespace."
        )

    sha256 = compute_text_sha256(clean_text)
    source_id = generate_id("src")
    
    doc = SourceDocument(
        source_id=source_id,
        filename=filename,
        media_type="text/plain",
        sha256=sha256,
        character_count=len(clean_text),
        page_count=None,
        extraction_status="success",
        warnings=[]
    )
    
    pages = [ParsedPage(page_number=None, text=clean_text, start_char=0, end_char=len(clean_text))]
    return ParsedResult(document=doc, pages=pages, full_text=clean_text)

def parse_pdf_bytes(content: bytes, filename: str) -> ParsedResult:
    """
    Extract text page-by-page from PDF bytes using pypdf.
    Page numbering is 1-based.
    """
    if not content:
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"PDF document '{filename}' is 0 bytes."
        )

    sha256 = compute_sha256(content)
    source_id = generate_id("src")
    warnings: List[str] = []

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"Failed to read PDF '{filename}': {str(exc)}"
        )

    pages: List[ParsedPage] = []
    full_text_parts: List[str] = []
    current_offset = 0

    page_count = len(reader.pages)
    if page_count == 0:
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"PDF '{filename}' contains 0 pages."
        )

    max_pages = 150
    pages_to_process = min(page_count, max_pages)
    if page_count > max_pages:
        warnings.append(f"Processed first {max_pages} pages of {page_count}-page PDF within safety processing limit.")

    for idx in range(pages_to_process):
        page = reader.pages[idx]
        page_number = idx + 1
        page_text = page.extract_text() or ""
        page_text = page_text.replace("\r\n", "\n").replace("\r", "\n")
        
        start_char = current_offset
        end_char = start_char + len(page_text)
        current_offset = end_char

        pages.append(ParsedPage(page_number=page_number, text=page_text, start_char=start_char, end_char=end_char))
        full_text_parts.append(page_text)

    full_text = "".join(full_text_parts)
    is_empty_or_scanned = len(full_text.strip()) < 20
    if is_empty_or_scanned:
        warnings.append("Scanned or image-only PDF detected. No extractable digital text layer found; OCR is required to extract text from scanned images.")

    doc = SourceDocument(
        source_id=source_id,
        filename=filename,
        media_type="application/pdf",
        sha256=sha256,
        character_count=len(full_text),
        page_count=page_count,
        extraction_status="scanned_pdf_no_text" if is_empty_or_scanned else "success",
        warnings=warnings
    )

    return ParsedResult(document=doc, pages=pages, full_text=full_text)

def parse_docx_bytes(content: bytes, filename: str) -> ParsedResult:
    """
    Extract text from DOCX bytes using python-docx.
    Page numbers are null because DOCX has flowing layout without fixed pages.
    """
    if not content:
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"DOCX document '{filename}' is 0 bytes."
        )

    sha256 = compute_sha256(content)
    source_id = generate_id("src")

    try:
        doc_obj = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"Failed to read DOCX '{filename}': {str(exc)}"
        )

    paragraphs_text: List[str] = []
    for p in doc_obj.paragraphs:
        if p.text:
            paragraphs_text.append(p.text)

    # Also extract text from tables
    for table in doc_obj.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs_text.append(" | ".join(row_texts))

    full_text = "\n\n".join(paragraphs_text).replace("\r\n", "\n").replace("\r", "\n")
    if not full_text.strip():
        raise SecondOpinionException(
            ErrorCode.SOURCE_EMPTY_OR_UNREADABLE,
            f"DOCX '{filename}' contains no extractable text."
        )

    doc = SourceDocument(
        source_id=source_id,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sha256=sha256,
        character_count=len(full_text),
        page_count=None,
        extraction_status="success",
        warnings=[]
    )

    pages = [ParsedPage(page_number=None, text=full_text, start_char=0, end_char=len(full_text))]
    return ParsedResult(document=doc, pages=pages, full_text=full_text)
