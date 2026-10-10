"""
Tests for document parsing (TXT, PDF, DOCX) and source chunking.
"""
import io
import pytest
from pypdf import PdfWriter
from docx import Document as DocxDocument

from backend.app.services.document_parser import (
    parse_plain_text,
    parse_pdf_bytes,
    parse_docx_bytes,
)
from backend.app.services.source_chunks import chunk_document
from backend.app.errors import SecondOpinionException, ErrorCode

def test_parse_plain_text():
    sample_text = "This is a reference document.\nIt contains multiple lines of text."
    result = parse_plain_text(sample_text, "test.txt")
    
    assert result.document.filename == "test.txt"
    assert result.document.media_type == "text/plain"
    assert result.document.character_count == len(sample_text)
    assert result.document.page_count is None
    assert len(result.document.sha256) == 64
    assert result.full_text == sample_text
    assert len(result.pages) == 1
    assert result.pages[0].page_number is None

def test_parse_plain_text_empty_raises():
    with pytest.raises(SecondOpinionException) as exc_info:
        parse_plain_text("   \n\t  ", "empty.txt")
    assert exc_info.value.detail["code"] == ErrorCode.SOURCE_EMPTY_OR_UNREADABLE.value

def test_parse_docx():
    # Build an in-memory docx
    doc = DocxDocument()
    doc.add_heading("Test Policy", level=1)
    doc.add_paragraph("This is paragraph one of the policy.")
    doc.add_paragraph("This is paragraph two.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Val 1"
    table.cell(1, 1).text = "Val 2"
    
    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()
    
    result = parse_docx_bytes(content, "policy.docx")
    assert result.document.filename == "policy.docx"
    assert "Test Policy" in result.full_text
    assert "paragraph one" in result.full_text
    assert "Header A | Header B" in result.full_text
    assert result.document.page_count is None

def test_parse_pdf():
    # Build a 2-page in-memory PDF using pypdf PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    
    buf = io.BytesIO()
    writer.write(buf)
    content = buf.getvalue()
    
    result = parse_pdf_bytes(content, "blank.pdf")
    assert result.document.filename == "blank.pdf"
    assert result.document.page_count == 2
    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert result.pages[1].page_number == 2

def test_chunking_offsets_match_text():
    text = " ".join([f"Word{i}" for i in range(300)])
    parsed = parse_plain_text(text, "words.txt")
    chunks = chunk_document(parsed, target_words=100, overlap_words=20)
    
    assert len(chunks) > 1
    for chunk in chunks:
        # Crucial invariant: sliced text from full_text MUST exactly match chunk.text
        slice_text = parsed.full_text[chunk.start_char:chunk.end_char]
        assert slice_text == chunk.text
        assert len(chunk.text) > 0
